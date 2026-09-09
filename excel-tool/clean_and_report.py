# -*- coding: utf-8 -*-
"""
clean_and_report.py

「input/売上データ_raw.csv」(表記ゆれだらけの汚いデータ) を読み込み、
  - 日付表記の統一 (YYYY-MM-DD)
  - 全角→半角変換
  - 前後空白の除去
  - 完全重複行・空行の除去
  - 商品名表記ゆれの統一 (マッピング辞書)
  - 単価テキスト ("1,200円" 等) の数値化
を行った上で、「output/売上レポート.xlsx」に4シート構成で書き出す。

    1. 整形済データ   … クリーニング後の明細 (Excelテーブル)
    2. 月別集計       … 月ごとの売上集計 + 棒グラフ
    3. 商品別集計     … 商品ごとの売上集計 (降順・合計行つき)
    4. クリーニングログ … 何件・何を直したかの記録

依存ライブラリ: openpyxl のみ (pandas 不使用、ポータビリティ重視)

実行方法:
    python clean_and_report.py
"""

import csv
import os
import re
import sys
import unicodedata
from collections import OrderedDict, defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

INPUT_FILE = os.path.join("input", "売上データ_raw.csv")
OUTPUT_FILE = os.path.join("output", "売上レポート.xlsx")

DEFAULT_YEAR = 2026  # 年省略の日付 ("1/5" 等) に補完する年

# ---------------------------------------------------------------------------
# 商品名の表記ゆれ統一マッピング
# (NFKC正規化 + 前後空白除去 + 全角括弧→半角括弧 まで済ませた文字列をキーにする)
# ---------------------------------------------------------------------------
PRODUCT_NAME_MAP = {
    "珈琲豆": "コーヒー豆",
    "紅茶": "紅茶葉",
    "こうちゃ葉": "紅茶葉",
    "ダンボール箱": "段ボール箱",
    "だんぼーる箱": "段ボール箱",
    "印刷用紙(A4)": "印刷用紙A4",
    "手袋(軍手)": "軍手",
    "せっちゃくざい": "接着剤",
    "杉材": "木材(杉)",
    "パン(食パン)": "食パン",
    "国産牛肉": "牛肉(国産)",
    "しょうゆ": "醤油",
    "コシヒカリ": "米(コシヒカリ)",
    "文具セット": "文房具セット",
}

DATE_HEADER = "日付"
COLUMNS = ["日付", "顧客名", "商品名", "数量", "単価", "担当者"]
OUTPUT_COLUMNS = ["日付", "顧客名", "商品名", "数量", "単価", "売上", "担当者"]


def has_fullwidth_digit(s: str) -> bool:
    return any(ch in "０１２３４５６７８９" for ch in s)


def normalize_text(s: str) -> str:
    """NFKC正規化 + 前後空白除去 + 全角括弧の半角化"""
    s = unicodedata.normalize("NFKC", s)
    s = s.strip()
    s = s.replace("（", "(").replace("）", ")")
    return s


DATE_PATTERNS = [
    ("slash_or_hyphen_with_year", re.compile(r"^(\d{4})[/\-](\d{1,2})[/\-](\d{1,2})$")),
    ("reiwa", re.compile(r"^令和(\d{1,2})年(\d{1,2})月(\d{1,2})日$")),
    ("no_year", re.compile(r"^(\d{1,2})[/\-](\d{1,2})$")),
]


def normalize_date(raw: str):
    """日付文字列を date オブジェクトに変換する。
    戻り値: (date または None, フォーマット種別 または None)
    """
    s = normalize_text(raw)

    m = DATE_PATTERNS[0][1].match(s)
    if m:
        y, mo, d = (int(x) for x in m.groups())
        try:
            return date(y, mo, d), "slash_or_hyphen_with_year"
        except ValueError:
            return None, None

    m = DATE_PATTERNS[1][1].match(s)
    if m:
        reiwa_y, mo, d = (int(x) for x in m.groups())
        y = reiwa_y + 2018  # 令和元年 = 2019年
        try:
            return date(y, mo, d), "reiwa"
        except ValueError:
            return None, None

    m = DATE_PATTERNS[2][1].match(s)
    if m:
        mo, d = (int(x) for x in m.groups())
        try:
            return date(DEFAULT_YEAR, mo, d), "no_year"
        except ValueError:
            return None, None

    return None, None


def normalize_product(raw: str, unmapped_counter: dict) -> str:
    key = normalize_text(raw)
    if key in PRODUCT_NAME_MAP:
        return PRODUCT_NAME_MAP[key]
    # NFKC正規化だけで正規名と一致するもの (例: ｺｰﾋｰ豆 -> コーヒー豆) はそのまま key を採用
    if key:
        return key
    unmapped_counter["count"] += 1
    return key


def parse_price(raw: str) -> int:
    s = normalize_text(raw)
    s = s.replace("円", "").replace("¥", "").replace(",", "")
    return int(s)


def parse_qty(raw: str) -> int:
    s = normalize_text(raw)
    return int(s)


def is_blank_row(row: dict) -> bool:
    return all((row.get(col) or "").strip() == "" for col in COLUMNS)


def had_leading_trailing_space(raw: str) -> bool:
    return raw != raw.strip()


def main():
    # Windowsコンソール(cp932)では ¥ 等のUnicode文字で落ちることがあるため、
    # 標準出力をUTF-8に切り替える (対応していない環境では無視する)
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    log = OrderedDict([
        ("読み込んだ生データ行数", 0),
        ("空白行の除去", 0),
        ("完全重複行の除去", 0),
        ("前後空白を除去したセル数", 0),
        ("全角数字→半角数字の変換件数", 0),
        ("日付を正規化した件数", 0),
        ("　- YYYY/M/D, YYYY-MM-DD形式", 0),
        ("　- 令和表記", 0),
        ("　- 年省略 (M/D)形式", 0),
        ("商品名の表記ゆれを統一した件数", 0),
        ("単価テキスト(円・¥・カンマ付き)を数値化した件数", 0),
        ("クリーニング後の最終行数", 0),
    ])

    with open(INPUT_FILE, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        raw_rows = list(reader)

    log["読み込んだ生データ行数"] = len(raw_rows)

    cleaned_rows = []
    seen = set()
    unmapped_counter = {"count": 0}

    for row in raw_rows:
        if is_blank_row(row):
            log["空白行の除去"] += 1
            continue

        # 前後空白カウント
        for col in COLUMNS:
            raw_val = row.get(col) or ""
            if had_leading_trailing_space(raw_val):
                log["前後空白を除去したセル数"] += 1
            if has_fullwidth_digit(raw_val):
                log["全角数字→半角数字の変換件数"] += 1

        d, date_fmt = normalize_date(row[DATE_HEADER])
        if d is None:
            # パースできない日付は行ごとスキップ (実務でも要目視確認対象)
            continue
        log["日付を正規化した件数"] += 1
        if date_fmt == "slash_or_hyphen_with_year":
            log["　- YYYY/M/D, YYYY-MM-DD形式"] += 1
        elif date_fmt == "reiwa":
            log["　- 令和表記"] += 1
        elif date_fmt == "no_year":
            log["　- 年省略 (M/D)形式"] += 1

        customer = normalize_text(row["顧客名"])
        product_raw = row["商品名"]
        product = normalize_product(product_raw, unmapped_counter)
        if normalize_text(product_raw) != product:
            log["商品名の表記ゆれを統一した件数"] += 1

        price_raw = row["単価"]
        if any(ch in price_raw for ch in ("円", "¥", ",")):
            log["単価テキスト(円・¥・カンマ付き)を数値化した件数"] += 1
        try:
            price = parse_price(price_raw)
            qty = parse_qty(row["数量"])
        except ValueError:
            continue

        staff = normalize_text(row["担当者"])

        key = (d.isoformat(), customer, product, qty, price, staff)
        if key in seen:
            log["完全重複行の除去"] += 1
            continue
        seen.add(key)

        cleaned_rows.append({
            "日付": d,
            "顧客名": customer,
            "商品名": product,
            "数量": qty,
            "単価": price,
            "売上": qty * price,
            "担当者": staff,
        })

    cleaned_rows.sort(key=lambda r: r["日付"])
    log["クリーニング後の最終行数"] = len(cleaned_rows)

    os.makedirs("output", exist_ok=True)
    write_report(cleaned_rows, log)

    print(f"[OK] {OUTPUT_FILE} を生成しました ({len(cleaned_rows)}行)")
    for k, v in log.items():
        print(f"  {k}: {v}")


# ---------------------------------------------------------------------------
# Excel 出力まわり
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill(start_color="FF2F5496", end_color="FF2F5496", fill_type="solid")
HEADER_FONT = Font(color="FFFFFFFF", bold=True, name="游ゴシック", size=11)
TITLE_FONT = Font(bold=True, size=14, name="游ゴシック", color="FF2F5496")
THIN_BORDER = Border(
    left=Side(style="thin", color="FFBFBFBF"),
    right=Side(style="thin", color="FFBFBFBF"),
    top=Side(style="thin", color="FFBFBFBF"),
    bottom=Side(style="thin", color="FFBFBFBF"),
)
TOTAL_FILL = PatternFill(start_color="FFDDEBF7", end_color="FFDDEBF7", fill_type="solid")


def autosize_columns(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def visual_width(s) -> int:
    s = str(s)
    total = 0
    for ch in s:
        total += 2 if unicodedata.east_asian_width(ch) in ("W", "F", "A") else 1
    return total


def write_seikei_sheet(wb, rows):
    ws = wb.active
    ws.title = "整形済データ"

    ws.append(OUTPUT_COLUMNS)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for r in rows:
        ws.append([r["日付"], r["顧客名"], r["商品名"], r["数量"], r["単価"], r["売上"], r["担当者"]])

    last_row = ws.max_row
    for row_cells in ws.iter_rows(min_row=2, max_row=last_row, max_col=len(OUTPUT_COLUMNS)):
        for cell in row_cells:
            cell.border = THIN_BORDER
        row_cells[0].number_format = "yyyy-mm-dd"
        row_cells[3].number_format = "#,##0"
        row_cells[4].number_format = "#,##0円"
        row_cells[5].number_format = "#,##0円"

    if last_row >= 2:
        table_ref = f"A1:G{last_row}"
        table = Table(displayName="SeikeiData", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False
        )
        ws.add_table(table)

    widths = []
    for col_idx, col_name in enumerate(OUTPUT_COLUMNS, start=1):
        max_w = visual_width(col_name)
        for row_cells in ws.iter_rows(min_row=2, max_row=last_row, min_col=col_idx, max_col=col_idx):
            val = row_cells[0].value
            if col_name == "日付" and val is not None:
                text = val.strftime("%Y-%m-%d")
            elif col_name in ("単価", "売上") and val is not None:
                text = f"{val:,}円"
            else:
                text = val
            max_w = max(max_w, visual_width(text))
        widths.append(max_w + 3)
    autosize_columns(ws, widths)
    ws.freeze_panes = "A2"


def write_monthly_sheet(wb, rows):
    ws = wb.create_sheet("月別集計")

    monthly = defaultdict(lambda: {"sales": 0, "count": 0})
    for r in rows:
        key = r["日付"].strftime("%Y-%m")
        monthly[key]["sales"] += r["売上"]
        monthly[key]["count"] += 1

    ws.append(["月", "売上合計", "件数"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for month_key in sorted(monthly.keys()):
        ws.append([month_key, monthly[month_key]["sales"], monthly[month_key]["count"]])

    last_row = ws.max_row
    for row_cells in ws.iter_rows(min_row=2, max_row=last_row, max_col=3):
        for cell in row_cells:
            cell.border = THIN_BORDER
        row_cells[1].number_format = "#,##0円"

    autosize_columns(ws, [12, 16, 10])
    ws.freeze_panes = "A2"

    if last_row >= 2:
        chart = BarChart()
        chart.type = "col"
        chart.title = "月別売上"
        chart.y_axis.title = "売上 (円)"
        chart.x_axis.title = "月"
        chart.style = 10

        data = Reference(ws, min_col=2, min_row=1, max_row=last_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 20
        chart.height = 10
        ws.add_chart(chart, f"E2")


def write_product_sheet(wb, rows):
    ws = wb.create_sheet("商品別集計")

    by_product = defaultdict(lambda: {"sales": 0, "count": 0})
    for r in rows:
        by_product[r["商品名"]]["sales"] += r["売上"]
        by_product[r["商品名"]]["count"] += 1

    ws.append(["商品名", "売上合計", "件数"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    ordered = sorted(by_product.items(), key=lambda kv: kv[1]["sales"], reverse=True)
    for name, agg in ordered:
        ws.append([name, agg["sales"], agg["count"]])

    total_sales = sum(agg["sales"] for _, agg in ordered)
    total_count = sum(agg["count"] for _, agg in ordered)
    ws.append(["合計", total_sales, total_count])

    last_row = ws.max_row
    for row_cells in ws.iter_rows(min_row=2, max_row=last_row, max_col=3):
        for cell in row_cells:
            cell.border = THIN_BORDER
        row_cells[1].number_format = "#,##0円"

    total_row = ws[last_row]
    for cell in total_row:
        cell.font = Font(bold=True)
        cell.fill = TOTAL_FILL

    max_name_w = max([visual_width("商品名")] + [visual_width(n) for n, _ in ordered] + [visual_width("合計")])
    autosize_columns(ws, [max_name_w + 3, 16, 10])
    ws.freeze_panes = "A2"


def write_log_sheet(wb, log):
    ws = wb.create_sheet("クリーニングログ")

    ws["A1"] = "データクリーニング処理ログ"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")

    ws.append([])
    ws.append(["項目", "件数"])
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for k, v in log.items():
        ws.append([k, v])

    last_row = ws.max_row
    for row_cells in ws.iter_rows(min_row=header_row + 1, max_row=last_row, max_col=2):
        for cell in row_cells:
            cell.border = THIN_BORDER
        row_cells[1].number_format = "#,##0"
        row_cells[1].alignment = Alignment(horizontal="right")

    max_item_w = max(visual_width(k) for k in list(log.keys()) + ["項目"])
    autosize_columns(ws, [max_item_w + 3, 12])


def write_report(rows, log):
    wb = Workbook()
    write_seikei_sheet(wb, rows)
    write_monthly_sheet(wb, rows)
    write_product_sheet(wb, rows)
    write_log_sheet(wb, log)
    wb.save(OUTPUT_FILE)


if __name__ == "__main__":
    main()
